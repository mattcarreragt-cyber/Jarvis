"""Extraction de texte multi-format — partagé par l'upload manuel et la sync Nextcloud.

Formats texte :  .txt .md
Documents     :  .pdf .docx .xlsx .pptx
Images        :  .png .jpg .jpeg .gif .webp .bmp
    → en mode Unraid (keyword), on indexe le nom + chemin + OCR (si tesseract dispo).
      La vraie description visuelle (caption) viendra d'un modèle vision sur Kubuntu
      (LLaVA via Ollama) — voir KUBUNTU_SETUP.md.

Les formats binaires legacy (.doc .xls .ppt) ne sont pas supportés nativement :
convertis-les en .docx/.xlsx/.pptx (Nextcloud/Collabora le fait automatiquement).
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger("jarvis.extract")

TEXT_EXT  = {".txt", ".md"}
PDF_EXT   = {".pdf"}
DOCX_EXT  = {".docx"}
XLSX_EXT  = {".xlsx"}
PPTX_EXT  = {".pptx"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
LEGACY_EXT = {".doc", ".xls", ".ppt"}

SUPPORTED_EXT = TEXT_EXT | PDF_EXT | DOCX_EXT | XLSX_EXT | PPTX_EXT | IMAGE_EXT


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXT


def kind_for(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in PDF_EXT:
        return "pdf"
    if ext in DOCX_EXT:
        return "docx"
    if ext in XLSX_EXT:
        return "xlsx"
    if ext in PPTX_EXT:
        return "pptx"
    return "doc"


# ─── Extracteurs par format ──────────────────────────────────────────────────

def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    return "\n".join(p.extract_text() or "" for p in reader.pages)


def _extract_docx(content: bytes) -> str:
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tableaux
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"# Feuille : {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_pptx(content: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(content))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# Diapositive {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = "".join(run.text for run in para.runs).strip()
                    if txt:
                        parts.append(txt)
    return "\n".join(parts)


def _ocr_image(content: bytes) -> str:
    """OCR optionnel via tesseract. Retourne '' si indisponible."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    try:
        img = Image.open(io.BytesIO(content))
        return pytesseract.image_to_string(img, lang="fra+eng").strip()
    except Exception as e:
        logger.debug("OCR indisponible/échoué : %s", e)
        return ""


def _extract_image(filename: str, content: bytes) -> str:
    """Contexte indexable d'une image : chemin + dossier + OCR éventuel.

    La description visuelle riche (caption) est ajoutée plus tard par le modèle
    vision sur Kubuntu — ici on garantit au moins la trouvabilité par nom/dossier.
    """
    path = Path(filename)
    folder = path.parent.name
    stem_words = path.stem.replace("_", " ").replace("-", " ")
    context = [
        f"Image : {path.name}",
        f"Dossier : {folder}" if folder else "",
        f"Mots du nom : {stem_words}",
    ]
    ocr = _ocr_image(content)
    if ocr:
        context.append(f"Texte détecté dans l'image :\n{ocr}")
    return "\n".join(c for c in context if c)


# ─── Point d'entrée ──────────────────────────────────────────────────────────

def extract_text(filename: str, content: bytes) -> str:
    """Extrait le texte indexable d'un fichier. Lève ValueError si non supporté."""
    ext = Path(filename).suffix.lower()
    try:
        if ext in TEXT_EXT:
            return content.decode("utf-8", errors="replace")
        if ext in PDF_EXT:
            return _extract_pdf(content)
        if ext in DOCX_EXT:
            return _extract_docx(content)
        if ext in XLSX_EXT:
            return _extract_xlsx(content)
        if ext in PPTX_EXT:
            return _extract_pptx(content)
        if ext in IMAGE_EXT:
            return _extract_image(filename, content)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Impossible de lire {ext} : {e}") from e

    if ext in LEGACY_EXT:
        raise ValueError(
            f"Format legacy {ext} non supporté — convertis en "
            f"{ext}x (Nextcloud/Collabora le fait automatiquement)."
        )
    raise ValueError(f"Format non supporté : {ext}")
